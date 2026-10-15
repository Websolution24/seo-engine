# src_utils.py
"""SEO_ENGINE — shared helpers.

Features:
- loads only the local .env on explicit request
- loguru logger -> logs/seo_engine.log + stdout
- TXT / JSON / CSV (legacy) helpers
- XLSX multi-sheet writer (overwrite or append)

Notes:
- The project uses this module name only: src_utils.py
"""

from __future__ import annotations

import csv
import json
import os
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from dotenv import load_dotenv
from loguru import logger


def find_project_root(start: Path) -> Path:
    """Find the project root (folder that contains input/ and output/)."""
    start = start.resolve()
    candidates = [start] + list(start.parents)[:5]
    for p in candidates:
        if (p / "input").exists() and (p / "output").exists():
            return p
    return start


PROJECT_ROOT = Path(__file__).resolve().parent


def iso_now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


# ENV (.env)
def load_project_env() -> None:
    """Load only the local project file, explicitly after paid-API consent."""
    env_path = PROJECT_ROOT / ".env"
    if env_path.is_file():
        load_dotenv(dotenv_path=env_path)


def get_env(var_name: str, default: Optional[str] = None) -> str:
    value = os.getenv(var_name)
    if value is None or value == "":
        if default is not None:
            return default
        logger.error(f"Missing environment variable: {var_name}")
        raise RuntimeError(f"Missing ENV var: {var_name}")
    return value


# LOGS
LOG_DIR = PROJECT_ROOT / "logs"
LOG_DIR.mkdir(parents=True, exist_ok=True)

logger.remove()
logger.add(
    LOG_DIR / "seo_engine.log",
    rotation="1 MB",
    level="INFO",
    format="{time:YYYY-MM-DD HH:mm:ss.SSS} | {level} | {name}:{function}:{line} - {message}",
)
logger.add(
    sys.stdout,
    level="INFO",
    format="{time:YYYY-MM-DD HH:mm:ss.SSS} | {level} | {message}",
)


# CSV (legacy)
def safe_spreadsheet_value(value: Any) -> Any:
    """Keep numbers unchanged; prevent untrusted strings becoming formulas."""
    if not isinstance(value, str):
        return value
    # XLSX forbids these C0 control characters; keep tab/newline/CR as text.
    cleaned = "".join(char for char in value if ord(char) >= 32 or char in "\t\r\n")
    significant = value.lstrip()
    while significant and (significant[0].isspace() or unicodedata.category(significant[0]).startswith("C")):
        significant = significant[1:]
    if significant.startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        return "'" + cleaned
    return cleaned


def read_csv(filepath: Path) -> List[Dict[str, str]]:
    filepath = filepath.resolve()
    logger.info(f"Reading CSV: {filepath}")
    if not filepath.exists():
        raise FileNotFoundError(str(filepath))
    with filepath.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        return list(reader)


def write_csv(filepath: Path, rows: List[Dict[str, Any]]) -> None:
    filepath.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        logger.warning(f"No rows to write: {filepath}")
        return
    with filepath.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writerow({column: safe_spreadsheet_value(column) for column in rows[0]})
        writer.writerows({column: safe_spreadsheet_value(value) for column, value in row.items()} for row in rows)


# TXT
def read_txt_lines(filepath: Path) -> List[str]:
    filepath = filepath.resolve()
    logger.info(f"Reading TXT: {filepath}")
    if not filepath.exists():
        raise FileNotFoundError(str(filepath))
    lines: List[str] = []
    for raw in filepath.read_text(encoding="utf-8").splitlines():
        s = raw.strip()
        if not s or s.startswith("#"):
            continue
        lines.append(s)
    return lines


def write_txt_lines(filepath: Path, lines: Iterable[str]) -> None:
    filepath.parent.mkdir(parents=True, exist_ok=True)
    data = "\n".join([str(x).strip() for x in lines if str(x).strip()])
    if data:
        data += "\n"
    filepath.write_text(data, encoding="utf-8")
    logger.info(f"TXT written: {filepath}")


def read_keywords_txt(filepath: Path) -> List[str]:
    """Alias: read lines and deduplicate while keeping order."""
    seen = set()
    out: List[str] = []
    for kw in read_txt_lines(filepath):
        if kw not in seen:
            seen.add(kw)
            out.append(kw)
    return out


# JSON
def read_json(filepath: Path) -> Any:
    filepath = filepath.resolve()
    logger.info(f"Reading JSON: {filepath}")
    if not filepath.exists():
        raise FileNotFoundError(str(filepath))
    with filepath.open(encoding="utf-8") as f:
        return json.load(f)


def write_json(filepath: Path, data: Any) -> None:
    filepath.parent.mkdir(parents=True, exist_ok=True)
    with filepath.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    logger.info(f"JSON written: {filepath}")


# XLSX
def write_xlsx_sheets(
    filepath: Path,
    sheets: Dict[str, List[Dict[str, Any]]],
    mode: str = "overwrite",
) -> None:
    """Write multiple sheets into an XLSX.

    sheets = {"SheetName": [ {col: value, ...}, ... ]}

    mode:
    - overwrite: create a brand new workbook
    - append: keep existing workbook, replace only provided sheets
    """
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter

    if mode not in {"overwrite", "append"}:
        raise ValueError("mode must be 'overwrite' or 'append'")

    filepath.parent.mkdir(parents=True, exist_ok=True)

    if mode == "append" and filepath.exists():
        wb = load_workbook(filepath)
    else:
        wb = Workbook()
        if wb.active:
            wb.remove(wb.active)

    for sheet_name, rows in sheets.items():
        safe_name = str(sheet_name)[:31]
        if safe_name in wb.sheetnames:
            wb.remove(wb[safe_name])
        ws = wb.create_sheet(title=safe_name)

        if not rows:
            ws.append(["(vide)"])
            continue

        columns: List[str] = []
        seen_cols = set()
        for r in rows:
            for k in r.keys():
                if k not in seen_cols:
                    seen_cols.add(k)
                    columns.append(k)

        ws.append([safe_spreadsheet_value(column) for column in columns])
        for r in rows:
            ws.append([safe_spreadsheet_value(r.get(c, "")) for c in columns])

        sample = rows[:200]
        for col_idx, col_name in enumerate(columns, start=1):
            max_len = len(str(col_name))
            for rr in sample:
                max_len = max(max_len, len(str(rr.get(col_name, ""))))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(60, max(10, max_len + 2))

    wb.save(filepath)
    logger.info(f"XLSX written: {filepath} (mode={mode})")
